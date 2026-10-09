"""Office files without LibreOffice, plus structure hints for either path.

render_native(): python-docx / python-pptx / openpyxl -> blocks per unit (document, slide, sheet part)
-> PDF with ReportLab Platypus. Layout is approximated, but reading order, heading sizes, tables and
images survive, which is what profiling, parsing and chunking need.

extract_hints(): things the PDF loses even with LibreOffice: speaker notes, slide titles and the
data behind native PowerPoint charts.

A unit is a list of blocks: ("h", level, text), ("p", text, indent), ("table", rows, caption) or
("img", bytes, max_height).
"""

import html
import io
import re
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Flowable, Image, PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

from .pdfgen import A4, cjk_font

SLIDE = (960, 540)  # 16:9 in points
A4_LANDSCAPE = (A4[1], A4[0])
MARGIN = 36

# Sizes in points: (body, h1, h2, h3, table cell)
DOC_SIZES = (10, 20, 16, 13, 9)
SLIDE_SIZES = (18, 30, 24, 20, 14)


def _e(s) -> str:
    if isinstance(s, float) and s.is_integer():
        s = int(s)
    return html.escape("" if s is None else str(s))


# ---------- docx ----------

def _docx_units(path: Path, with_images: bool = True) -> tuple[list[list], dict]:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    d = docx.Document(str(path))
    blocks, headings, n_img = [], [], 0
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = DocxParagraph(el, d)
            style = (p.style.name if p.style is not None else "") or ""
            text = p.text.strip()
            m = re.match(r"(?:Heading|제목)\s*(\d)", style)
            level = 1 if style in ("Title", "제목") else int(m.group(1)) if m else 0
            for blip in el.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                if rid and rid in d.part.related_parts:
                    n_img += 1
                    if with_images:
                        blocks.append(("img", d.part.related_parts[rid].blob, None))
            if not text:
                continue
            if level:
                headings.append({"level": level, "text": text})
                blocks.append(("h", min(level, 3), text))
            else:
                blocks.append(("p", text, 0))
        elif el.tag == qn("w:tbl"):
            t = DocxTable(el, d)
            blocks.append(("table", [[c.text.strip() for c in r.cells] for r in t.rows], None))
    return [blocks], {"kind": "docx", "headings": headings, "images": n_img}


# ---------- pptx ----------

def _chart_rows(chart) -> tuple[str, list[list]]:
    title = chart.chart_title.text_frame.text if chart.has_title and chart.chart_title.has_text_frame else ""
    plot = chart.plots[0]
    cats = [str(c) for c in plot.categories]
    series = list(plot.series)
    rows = [["항목"] + [s.name for s in series]]
    for i, c in enumerate(cats):
        rows.append([c] + [s.values[i] if i < len(s.values) else "" for s in series])
    return title, rows


def _shapes(shapes):
    """Flatten group shapes, in visual order (top, then left)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    flat = []
    for s in shapes:
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            flat.extend(_shapes(s.shapes))
        else:
            flat.append(s)
    return sorted(flat, key=lambda s: ((s.top or 0), (s.left or 0)))


def _pptx_slides(path: Path, with_images: bool = True) -> tuple[list[list], dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    units, slides = [], []
    for i, slide in enumerate(prs.slides):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        body, charts = [], []
        if title:
            body.append(("h", 1, title))
        for s in _shapes(slide.shapes):
            if title_shape is not None and s.shape_id == title_shape.shape_id:
                continue
            if getattr(s, "has_chart", False) and s.has_chart:
                ctitle, rows = _chart_rows(s.chart)
                charts.append({"title": ctitle, "rows": rows})
                body.append(("table", rows, f"차트: {ctitle}" if ctitle else "차트"))
            elif getattr(s, "has_table", False) and s.has_table:
                body.append(("table", [[c.text.strip() for c in r.cells] for r in s.table.rows], None))
            elif s.shape_type == MSO_SHAPE_TYPE.PICTURE and with_images:
                body.append(("img", s.image.blob, 260))
            elif s.has_text_frame:
                for para in s.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        body.append(("p", t, para.level or 0))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        units.append(body)
        slides.append({"slide": i, "title": title, "notes": notes, "charts": charts})
    return units, {"kind": "pptx", "slides": slides}


# ---------- xlsx ----------

def _xlsx_units(path: Path, rows_per_page: int, max_rows: int) -> tuple[list[list], dict]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    units, sheets = [], []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            if len(rows) >= max_rows + 1:
                break
            rows.append(["" if v is None else v for v in r])
        while rows and not any(str(v).strip() for v in rows[-1]):
            rows.pop()
        width = max((i + 1 for r in rows for i, v in enumerate(r) if str(v).strip()), default=0)
        rows = [r[:width] for r in rows]
        truncated = ws.max_row is not None and ws.max_row > max_rows + 1
        sheets.append({"sheet": ws.title, "rows": max(0, len(rows) - 1), "cols": width, "truncated": truncated})
        if not rows:
            continue
        header, data = rows[0], rows[1:] or [[]]
        # One unit per block of rows, header repeated, so every page (and chunk) is a self-describing table.
        for start in range(0, len(data), rows_per_page):
            block = data[start:start + rows_per_page]
            label = f"{ws.title} (행 {start + 2}–{start + 1 + len(block)})"
            units.append([("h", 2, label), ("table", [header] + block, None)])
    wb.close()
    return units, {"kind": "xlsx", "sheets": sheets}


# ---------- render ----------

class _UnitStart(Flowable):
    """Zero-size marker that records the 0-based page its unit starts on."""

    def __init__(self, starts: list[int]):
        super().__init__()
        self.starts = starts

    def wrap(self, *_):
        return 0, 0

    def draw(self) -> None:
        self.starts.append(self.canv.getPageNumber() - 1)


def _styles(font: str, sizes: tuple) -> dict:
    body, h1, h2, h3, cell = sizes
    return {
        "p": ParagraphStyle("p", fontName=font, fontSize=body, leading=body * 1.45, spaceAfter=body * 0.5),
        1: ParagraphStyle("h1", fontName=font, fontSize=h1, leading=h1 * 1.25, spaceAfter=h1 * 0.45),
        2: ParagraphStyle("h2", fontName=font, fontSize=h2, leading=h2 * 1.25, spaceBefore=h2 * 0.6, spaceAfter=h2 * 0.4),
        3: ParagraphStyle("h3", fontName=font, fontSize=h3, leading=h3 * 1.25, spaceBefore=h3 * 0.6, spaceAfter=h3 * 0.3),
        "cell": ParagraphStyle("cell", fontName=font, fontSize=cell, leading=cell * 1.3),
    }


def _table(rows: list[list], caption: str | None, st: dict, width: float) -> list:
    ncols = max((len(r) for r in rows), default=0)
    if not ncols:
        return []
    data = [[Paragraph(_e(c), st["cell"]) for c in list(r) + [""] * (ncols - len(r))] for r in rows]
    t = Table(data, colWidths=[width / ncols] * ncols, repeatRows=1, hAlign="LEFT", spaceBefore=6, spaceAfter=6)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.6, colors.HexColor("#555555")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    return ([Paragraph(_e(caption), st["cell"])] if caption else []) + [t]


def _image(blob: bytes, max_h: float | None, width: float, height: float) -> list:
    try:
        img = Image(io.BytesIO(blob))
    except Exception:  # formats Pillow cannot read (EMF/WMF, ...): the page keeps its text
        return []
    scale = min(1.0, width / img.imageWidth, (max_h or height) / img.imageHeight)
    img.drawWidth, img.drawHeight = img.imageWidth * scale, img.imageHeight * scale
    img.hAlign = "LEFT"
    return [img]


def _write(units: list[list], out: Path, page: tuple[float, float], sizes: tuple, font: str) -> list[int]:
    """Each unit starts on a new page and may flow onto more. Returns the first page index of every unit."""
    st = _styles(font, sizes)
    width, height = page[0] - 2 * MARGIN, page[1] - 2 * MARGIN
    starts: list[int] = []
    story: list = []
    for n, unit in enumerate(units):
        if n:
            story.append(PageBreak())
        story.append(_UnitStart(starts))
        for b in unit:
            if b[0] == "h":
                story.append(Paragraph(_e(b[2]), st[b[1]]))
            elif b[0] == "p":
                story.append(Paragraph(_e(b[1]), st["p"].clone("pi", leftIndent=b[2] * st["p"].fontSize * 2) if b[2] else st["p"]))
            elif b[0] == "table":
                story.extend(_table(b[1], b[2], st, width))
            elif b[0] == "img":
                story.extend(_image(b[1], b[2], width, height))
    doc = SimpleDocTemplate(str(out), pagesize=page, leftMargin=MARGIN, rightMargin=MARGIN, topMargin=MARGIN,
                            bottomMargin=MARGIN, title=out.stem)
    doc.build(story)
    return starts


def render_native(src: Path, pdf: Path, fmt: str, xlsx_rows_per_page: int = 24, xlsx_max_rows: int = 5000,
                  font_path: str = "") -> dict:
    """Writes the PDF `pdf`; returns the hints."""
    pdf.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx":
        units, hints = _docx_units(src)
        page, sizes = A4, DOC_SIZES
    elif fmt == "pptx":
        units, hints = _pptx_slides(src)
        page, sizes = SLIDE, SLIDE_SIZES
    elif fmt == "xlsx":
        units, hints = _xlsx_units(src, xlsx_rows_per_page, xlsx_max_rows)
        page, sizes = A4_LANDSCAPE, DOC_SIZES
    else:
        raise ValueError(f"no native renderer for .{fmt}")
    font = cjk_font(font_path)
    starts = _write(units or [[]], pdf, page, sizes, font["name"])
    if hints["kind"] == "pptx":
        for s, first in zip(hints["slides"], starts):
            s["page"] = first
    hints["unit_pages"] = starts
    hints["font"] = {"source": font["source"], "embedded": font["embedded"], "rule_id": font["rule_id"]}
    return hints


def extract_hints(src: Path, fmt: str) -> dict:
    """Hints for a LibreOffice-converted file (LibreOffice renders one PDF page per slide)."""
    if fmt == "pptx":
        _, hints = _pptx_slides(src, with_images=False)
        for s in hints["slides"]:
            s["page"] = s["slide"]
        return hints
    if fmt == "docx":
        _, hints = _docx_units(src, with_images=False)
        return hints
    return {}
