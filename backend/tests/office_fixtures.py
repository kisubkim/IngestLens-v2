"""Generate small Office files with known structure for tests."""

import io
from pathlib import Path

from PIL import Image  # comes with python-pptx


def _png(w: int = 120, h: int = 80) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (120, 160, 220)).save(buf, "PNG")
    return buf.getvalue()


def make_docx(path: Path) -> Path:
    import docx

    d = docx.Document()
    d.add_heading("장비 운영 매뉴얼", level=0)
    for i, sec in enumerate(["개요", "설치 절차", "점검 항목"], start=1):
        d.add_heading(f"{i}. {sec}", level=1)
        d.add_paragraph(f"{sec}에 대한 설명입니다. " * 25)
        if i == 2:
            d.add_heading("2.1 준비물", level=2)
            d.add_paragraph("전원 케이블, 통신 케이블, 설정 파일을 준비한다. " * 6)
            d.add_picture(io.BytesIO(_png()))
        if i == 3:
            t = d.add_table(rows=4, cols=3)
            for r, row in enumerate([["항목", "주기", "기준"], ["온도", "매일", "25±2"], ["압력", "주간", "1.2"], ["진동", "월간", "0.5"]]):
                for c, v in enumerate(row):
                    t.cell(r, c).text = v
    d.save(str(path))
    return path


def make_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "프로젝트 개요"
    s.placeholders[1].text_frame.text = "문서 자동 분석 파이프라인"
    p = s.placeholders[1].text_frame.add_paragraph()
    p.text, p.level = "로컬 vLLM 사용", 1
    s.notes_slide.notes_text_frame.text = "발표자 노트: 이 슬라이드에서는 배경을 설명한다."

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "월별 생산량"
    data = CategoryChartData()
    data.categories = ["1월", "2월", "3월"]
    data.add_series("생산량", (120, 135, 150))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(4.5), data)

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "점검 결과"
    t = s.shapes.add_table(3, 2, Inches(1), Inches(1.5), Inches(6), Inches(1.5)).table
    for r, row in enumerate([["항목", "결과"], ["온도", "정상"], ["압력", "주의"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    s.shapes.add_picture(io.BytesIO(_png()), Inches(8), Inches(1.5))
    prs.save(str(path))
    return path


def make_xlsx(path: Path, rows: int = 130) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "측정값"
    ws.append(["일자", "장비", "온도", "압력"])
    for i in range(rows):
        ws.append([f"2026-09-{i % 30 + 1:02d}", f"EQ-{i % 7}", 20 + i % 10, round(1 + (i % 5) * 0.1, 1)])
    ws2 = wb.create_sheet("요약")
    ws2.append(["장비", "평균 온도"])
    ws2.append(["EQ-0", 24.5])
    wb.save(str(path))
    return path
